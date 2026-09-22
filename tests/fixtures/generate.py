"""Builds the synthetic fixture files used by the parser tests.

Authored in code so ground truth is exact, and deterministic so that
content-addressed document ids stay stable across runs.

Run directly to regenerate: `uv run python -m tests.fixtures.generate`
"""

from __future__ import annotations

import re
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from PIL import Image, ImageDraw
from reportlab import rl_config
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas as rl_canvas

__all__ = ["build_all", "FIXTURE_DIR", "PAGE_WIDTH", "PAGE_HEIGHT"]

FIXTURE_DIR = Path(__file__).parent / "files"
PAGE_WIDTH, PAGE_HEIGHT = LETTER  # 612 x 792 points

BODY_LINES = [
    "The Payment API is responsible for processing payment requests.",
    "It validates each request, reserves funds, and writes an entry to",
    "the ledger before acknowledging the caller.",
]

DEPENDENCY_LINES = [
    "The Payment API depends on Redis for caching and on PostgreSQL",
    "for persistent storage. The Payments Team owns the service.",
]

# Three spaced paragraphs per column: consecutive lines would be merged into a
# single block by the PDF text extractor, leaving too few blocks for column
# detection to have anything to work with.
LEFT_COLUMN = [
    ["Incident INC-2391 affected the", "Payment API in production."],
    ["Redis connection timeouts rose", "sharply during peak traffic."],
    ["The Payments Team resolved it", "by resizing the pool."],
]

RIGHT_COLUMN = [
    ["The Gateway routes requests to", "the Payment API over gRPC."],
    ["PostgreSQL holds the ledger and", "the settlement records."],
    ["The Platform Team owns the", "Gateway and its runbook."],
]


def _canvas(path: Path) -> rl_canvas.Canvas:
    rl_config.invariant = 1
    return rl_canvas.Canvas(str(path), pagesize=LETTER, invariant=1)


def build_diagram_png(path: Path) -> Path:
    """A deterministic architecture diagram: boxes joined by arrows."""
    image = Image.new("RGB", (480, 300), "white")
    draw = ImageDraw.Draw(image)

    boxes = {
        "API Gateway": (160, 20, 320, 70),
        "Payment Service": (160, 115, 320, 165),
        "Redis": (30, 220, 180, 270),
        "PostgreSQL": (300, 220, 450, 270),
    }
    for label, (x0, y0, x1, y1) in boxes.items():
        draw.rectangle((x0, y0, x1, y1), outline="black", width=2)
        draw.text((x0 + 12, y0 + 18), label, fill="black")

    draw.line((240, 70, 240, 115), fill="black", width=2)
    draw.line((200, 165, 105, 220), fill="black", width=2)
    draw.line((280, 165, 375, 220), fill="black", width=2)

    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, format="PNG", optimize=False)
    return path


def build_simple(path: Path) -> Path:
    """One page: a title, a body paragraph, a subheading, and a running footer."""
    canvas = _canvas(path)

    canvas.setFont("Helvetica-Bold", 18)
    canvas.drawString(72, 720, "Payment Architecture")

    canvas.setFont("Helvetica", 11)
    y = 686
    for line in BODY_LINES:
        canvas.drawString(72, y, line)
        y -= 14

    canvas.setFont("Helvetica-Bold", 14)
    canvas.drawString(72, y - 18, "Dependencies")

    canvas.setFont("Helvetica", 11)
    y -= 46
    for line in DEPENDENCY_LINES:
        canvas.drawString(72, y, line)
        y -= 14

    canvas.setFont("Helvetica", 8)
    canvas.drawCentredString(PAGE_WIDTH / 2, 40, "Knowledge OS Internal")

    canvas.showPage()
    canvas.save()
    return path


def build_two_column(path: Path) -> Path:
    """Two pages of two-column body text with a running header and page numbers."""
    canvas = _canvas(path)

    for page_number in (1, 2):
        canvas.setFont("Helvetica", 9)
        canvas.drawString(72, 750, "Payment Platform Handbook")

        canvas.setFont("Helvetica-Bold", 16)
        canvas.drawString(72, 706, f"Section {page_number}")

        canvas.setFont("Helvetica", 11)
        for column_x, paragraphs in ((72, LEFT_COLUMN), (340, RIGHT_COLUMN)):
            y = 670
            for paragraph in paragraphs:
                for line in paragraph:
                    canvas.drawString(column_x, y, line)
                    y -= 14
                y -= 28  # paragraph gap, wide enough to break the text block

        canvas.setFont("Helvetica", 9)
        canvas.drawCentredString(PAGE_WIDTH / 2, 40, str(page_number))
        canvas.showPage()

    canvas.save()
    return path


def build_table_figure(path: Path, diagram: Path) -> Path:
    """One page: heading, a ruled table, and a captioned diagram."""
    canvas = _canvas(path)

    canvas.setFont("Helvetica-Bold", 18)
    canvas.drawString(72, 720, "Service Inventory")

    rows = [
        ["Service", "Datastore", "Owner"],
        ["Payment API", "PostgreSQL", "Payments Team"],
        ["Ledger", "Redis", "Payments Team"],
        ["Gateway", "None", "Platform Team"],
    ]
    col_x = [72, 232, 372, 540]
    top = 690
    row_height = 22

    for row_index, row in enumerate(rows):
        y = top - row_index * row_height
        canvas.setFont("Helvetica-Bold" if row_index == 0 else "Helvetica", 10)
        for col_index, cell in enumerate(row):
            canvas.drawString(col_x[col_index] + 4, y - 15, cell)
        canvas.line(col_x[0], y, col_x[-1], y)

    bottom = top - len(rows) * row_height
    canvas.line(col_x[0], bottom, col_x[-1], bottom)
    for x in col_x:
        canvas.line(x, top, x, bottom)

    canvas.drawImage(ImageReader(str(diagram)), 120, 300, width=360, height=225, mask=None)

    canvas.setFont("Helvetica", 9)
    canvas.drawString(120, 286, "Figure 1: Payment service architecture")

    canvas.showPage()
    canvas.save()
    return path


# A fixed instant for workbook metadata, and a fixed zip entry timestamp. Both
# default to "now", which would make the .xlsx bytes - and therefore its
# content-addressed document id - drift between sessions.
FIXED_TIME = datetime(2024, 1, 1, 0, 0, 0)
FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
FIXED_STAMP = b"2024-01-01T00:00:00Z"
CORE_PROPERTIES = "docProps/core.xml"

# openpyxl honours `properties.created` but stamps `dcterms:modified` with the
# wall clock at save time, so the same workbook differs between runs.
_MODIFIED = re.compile(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)")


def _normalize_ooxml(path: Path) -> None:
    """Rewrite an OOXML zip with sorted entries and every timestamp pinned."""
    with zipfile.ZipFile(path) as source:
        entries = sorted(
            ((info.filename, source.read(info.filename)) for info in source.infolist()),
            key=lambda pair: pair[0],
        )

    temporary = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as target:
        for name, data in entries:
            if name == CORE_PROPERTIES:
                data = _MODIFIED.sub(rb"\g<1>" + FIXED_STAMP + rb"\g<2>", data)
            info = zipfile.ZipInfo(name, date_time=FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(info, data)
    temporary.replace(path)


def build_sheet(path: Path) -> Path:
    """A two-sheet workbook with headers, a merged title and a formula."""
    workbook = Workbook()
    workbook.properties.created = FIXED_TIME
    workbook.properties.modified = FIXED_TIME

    services = workbook.active
    services.title = "Services"
    services["A1"] = "Service Inventory"
    services.merge_cells("A1:C1")
    services.append([])
    services.append(["Service", "Datastore", "Owner"])
    services.append(["Payment API", "PostgreSQL", "Payments Team"])
    services.append(["Ledger", "Redis", "Payments Team"])

    incidents = workbook.create_sheet("Incidents")
    incidents.append(["Incident", "Service", "Hours"])
    incidents.append(["INC-2391", "Payment API", 4])
    incidents.append(["INC-2402", "Ledger", 2])
    incidents["C4"] = "=SUM(C2:C3)"

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    _normalize_ooxml(path)
    return path


def build_all(out_dir: Path | None = None) -> dict[str, Path]:
    target = Path(out_dir) if out_dir is not None else FIXTURE_DIR
    target.mkdir(parents=True, exist_ok=True)

    diagram = build_diagram_png(target / "diagram.png")
    return {
        "simple": build_simple(target / "simple.pdf"),
        "two_column": build_two_column(target / "two_column.pdf"),
        "table_figure": build_table_figure(target / "table_figure.pdf", diagram),
        "diagram_png": diagram,
        "sheet": build_sheet(target / "inventory.xlsx"),
    }


if __name__ == "__main__":
    for name, built in build_all().items():
        print(f"{name}: {built}")
