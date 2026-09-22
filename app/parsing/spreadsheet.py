from __future__ import annotations

from pathlib import Path
from typing import Any, ClassVar

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.ir.ids import element_id
from app.ir.model import BBox, Document, Element, ElementType, Page, TableData
from app.ir.tree import assign_parents
from app.parsing.base import BlobSink, ParseError, source_identity

__all__ = ["SpreadsheetParser", "CELL_WIDTH", "ROW_HEIGHT"]

CELL_WIDTH = 80.0
ROW_HEIGHT = 18.0
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _stringify(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


class SpreadsheetParser:
    """One page per worksheet: a HEADING for the sheet name, then one TABLE.

    Spreadsheets have no page geometry, so bboxes are synthetic grid
    coordinates and every element is tagged `synthetic_bbox`.
    """

    extensions: ClassVar[frozenset[str]] = frozenset({".xlsx"})
    version: ClassVar[str] = "1"

    def parse(self, path: Path, *, blobs: BlobSink) -> Document:
        path = Path(path)
        short_id, checksum = source_identity(path)

        try:
            values_book = load_workbook(path, data_only=True, read_only=False)
            formula_book = load_workbook(path, data_only=False, read_only=False)
        except Exception as exc:
            raise ParseError(f"cannot open {path} as a workbook: {exc}") from exc

        document = Document(
            id=short_id,
            source_name=path.name,
            mime=XLSX_MIME,
            checksum=checksum,
            title=path.stem,
            meta={
                "parser": "spreadsheet",
                "parser_version": self.version,
                "sheet_names": list(values_book.sheetnames),
            },
        )

        order = 0
        for page_number, sheet_name in enumerate(values_book.sheetnames, start=1):
            grid = self._grid(values_book[sheet_name])
            formulas = self._formulas(formula_book[sheet_name])
            table, banner = self._table(grid)

            n_cols = max(table.n_cols, len(table.headers), 1)
            n_rows = max(table.n_rows, 1)

            heading = Element(
                id=element_id(page_number, 0),
                type=ElementType.HEADING,
                page_number=page_number,
                bbox=BBox(x0=0, y0=0, x1=CELL_WIDTH * n_cols, y1=ROW_HEIGHT),
                order=order,
                level=1,
                text=sheet_name,
                attrs={"synthetic_bbox": True},
            )
            order += 1

            table_element = Element(
                id=element_id(page_number, 1),
                type=ElementType.TABLE,
                page_number=page_number,
                bbox=BBox(
                    x0=0,
                    y0=ROW_HEIGHT,
                    x1=CELL_WIDTH * n_cols,
                    y1=ROW_HEIGHT * (n_rows + 2),
                ),
                order=order,
                table=table,
                attrs={
                    "synthetic_bbox": True,
                    "sheet": sheet_name,
                    "formulas": formulas,
                    "table_title": banner,
                },
            )
            order += 1

            document.pages.append(
                Page(
                    number=page_number,
                    width=CELL_WIDTH * n_cols,
                    height=ROW_HEIGHT * (n_rows + 3),
                    elements=[heading, table_element],
                )
            )

        assign_parents([e for page in document.pages for e in page.elements])
        document.meta["page_count"] = len(document.pages)
        return document

    def _grid(self, sheet: Any) -> list[list[str]]:
        """Rows of stringified values, with merged ranges filled across."""
        rows = [
            [_stringify(cell) for cell in row] for row in sheet.iter_rows(values_only=True)
        ]

        for merged in sorted(str(r) for r in sheet.merged_cells.ranges):
            bounds = sheet[merged]
            value = _stringify(bounds[0][0].value)
            if not value:
                continue
            for row in bounds:
                for cell in row:
                    r, c = cell.row - 1, cell.column - 1
                    if r < len(rows) and c < len(rows[r]):
                        rows[r][c] = value

        while rows and not any(cell for cell in rows[-1]):
            rows.pop()
        return rows

    def _formulas(self, sheet: Any) -> dict[str, str]:
        found: dict[str, str] = {}
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    found[f"{get_column_letter(cell.column)}{cell.row}"] = cell.value
        return dict(sorted(found.items()))

    def _table(self, grid: list[list[str]]) -> tuple[TableData, str | None]:
        """Split a grid into a title banner, a header row, and data rows."""
        rows = [row for row in grid if any(cell for cell in row)]
        if not rows:
            return TableData(), None

        n_cols = max(len(row) for row in rows)
        padded = [row + [""] * (n_cols - len(row)) for row in rows]

        # A merged full-width row is a title banner. Left in place it would be
        # mistaken for the header row, burying the real column names in the data.
        banner: str | None = None
        if n_cols > 1 and padded[0][0] and len(set(padded[0])) == 1:
            banner = padded[0][0]
            padded = padded[1:]

        headers: list[str] = []
        for index, row in enumerate(padded):
            if all(cell for cell in row):
                headers = row
                padded = padded[index + 1 :]
                break

        return (
            TableData(headers=headers, rows=padded, n_rows=len(padded), n_cols=n_cols),
            banner,
        )
